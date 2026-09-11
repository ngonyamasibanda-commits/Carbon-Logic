#!/usr/bin/env python3
"""Extract every Open CEDA UK GHG_t_Raw sector into src/lib/ceda-data.ts.

The xlsx is not committed; re-run when Watershed publishes a new CEDA year.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

XLSX = Path("/home/ubuntu/.cursor/projects/workspace/uploads/Open_CEDA_by_Watershed_9607.xlsx")
OUT = Path("/workspace/src/lib/ceda-data.ts")

NAICS_GROUPS = [
    ("11", "Agriculture, forestry, fishing, and hunting"),
    ("21", "Mining"),
    ("22", "Utilities"),
    ("23", "Construction"),
    ("31", "Manufacturing"),
    ("32", "Manufacturing"),
    ("33", "Manufacturing"),
    ("42", "Wholesale trade"),
    ("44", "Retail trade"),
    ("45", "Retail trade"),
    ("48", "Transportation and warehousing"),
    ("49", "Transportation and warehousing"),
    ("51", "Information"),
    ("52", "Finance and insurance"),
    ("53", "Real estate and rental"),
    ("54", "Professional, scientific, and technical services"),
    ("55", "Management of companies"),
    ("56", "Administrative, support, and waste services"),
    ("61", "Educational services"),
    ("62", "Health care and social assistance"),
    ("71", "Arts, entertainment, and recreation"),
    ("72", "Accommodation and food services"),
    ("81", "Other services"),
    ("92", "Public administration"),
]


def group_for(code: str) -> str:
    digits = "".join(ch for ch in code if ch.isdigit())
    prefix = digits[:2]
    for key, label in NAICS_GROUPS:
        if prefix == key:
            return label
    return "Other / special industries"


def fmt(n: float) -> str:
    s = f"{n:.12f}".rstrip("0").rstrip(".")
    return s


def code_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def main() -> None:
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    raw = wb["GHG_t_Raw"]
    names: list | None = None
    codes: list | None = None
    gbr: list | None = None
    for i, row in enumerate(raw.iter_rows(min_row=3, values_only=True), 3):
        vals = list(row)
        if i == 3:
            names = vals
        elif i == 4:
            codes = vals
        elif vals and vals[0] == "GBR":
            gbr = vals
            break
    if not names or not codes or not gbr:
        raise SystemExit("GBR row missing from GHG_t_Raw")

    pp = wb["Purchaser - producer conversion"]
    pp_codes = pp_vals = None
    for i, row in enumerate(pp.iter_rows(min_row=5, max_row=6, values_only=True), 5):
        if i == 5:
            pp_codes = list(row)
        else:
            pp_vals = list(row)
    pi = wb["Sector level Price Index"]
    pi_codes = pi_2025 = None
    for row in pi.iter_rows(min_row=5, max_row=13, values_only=True):
        vals = list(row)
        if vals[0] == "Sector Code":
            pi_codes = vals
        elif str(vals[0]) == "2025":
            pi_2025 = vals
    wb.close()
    if not pp_codes or not pp_vals or not pi_codes or not pi_2025:
        raise SystemExit("purchaser-producer or 2025 price index missing")

    pp_map = {code_str(c): float(v) for c, v in zip(pp_codes, pp_vals) if code_str(c) and isinstance(v, (int, float))}
    pi_map = {code_str(c): float(v) for c, v in zip(pi_codes, pi_2025) if code_str(c) and isinstance(v, (int, float))}

    rows = []
    for name, code, kg in zip(names[3:], codes[3:], gbr[3:]):
        code_s = code_str(code)
        if not code_s or not isinstance(kg, (int, float)):
            continue
        pp_v = pp_map.get(code_s)
        pi_v = pi_map.get(code_s)
        if pp_v is None or pi_v is None:
            raise SystemExit(f"missing conversion for {code_s}")
        key = f"ceda_gbr_{code_s.lower()}_usd"
        rows.append(
            {
                "key": key,
                "code": code_s,
                "name": str(name).replace("\n", " ").strip(),
                "kgPerUsd": float(kg),
                "purchaserProducer": float(pp_v),
                "priceIndex2025": float(pi_v),
                "group": group_for(code_s),
            }
        )

    lines = [
        "/** Open CEDA 2025 UK GHG_t_Raw sectors. Generated — do not edit by hand. */",
        "",
        "export type CedaSectorRow = {",
        "  key: string",
        "  code: string",
        "  name: string",
        "  kgPerUsd: number",
        "  purchaserProducer: number",
        "  priceIndex2025: number",
        "  group: string",
        "}",
        "",
        "export const CEDA_GBR_SECTORS = [",
    ]
    for row in rows:
        lines.append(
            "  { "
            + f'key: {row["key"]!r}, code: {row["code"]!r}, name: {row["name"]!r}, '
            + f'kgPerUsd: {fmt(row["kgPerUsd"])}, purchaserProducer: {fmt(row["purchaserProducer"])}, '
            + f'priceIndex2025: {fmt(row["priceIndex2025"])}, group: {row["group"]!r}'
            + " },"
        )
    lines.append("] as const")
    lines.append("")
    lines.append("export type CedaSector = (typeof CEDA_GBR_SECTORS)[number]")
    lines.append("")
    OUT.write_text("\n".join(lines) + "\n")
    print(f"wrote {len(rows)} CEDA UK sectors to {OUT}")


if __name__ == "__main__":
    main()
