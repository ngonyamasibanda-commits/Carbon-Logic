#!/usr/bin/env python3
"""Extract EPA Hub 2026 published types into src/lib/epa-hub-data.ts.

eGRID, GWP, freight, waste, and travel numbers come from the 2026 Hub PDF.
Stationary combustion physical factors come from the 2025 Hub workbook
(Table 1 kg CO₂ / g CH₄ / g N₂O per unit are unchanged in 2026) and are
converted to kg CO₂e with the 2026 AR6 GWPs (CH₄ 27, N₂O 273).
Workbooks/PDFs are not committed.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path("/home/ubuntu/.cursor/projects/workspace/uploads/ghg-emission-factors-hub-2025_27ef.xlsx")
OUT = Path("/workspace/src/lib/epa-hub-data.ts")

GWP_CH4 = 27
GWP_N2O = 273

# EPA Hub 2026 Table 6 — Total Output (not non-baseload), eGRID2024.
EGRID = [
    ("AKGD", "ASCC Alaska Grid", 914.7, 0.092, 0.013, 0.044),
    ("AKMS", "ASCC Miscellaneous", 493.8, 0.024, 0.004, 0.044),
    ("AZNM", "WECC Southwest", 665.8, 0.035, 0.005, 0.044),
    ("CAMX", "WECC California", 403.6, 0.025, 0.003, 0.044),
    ("ERCT", "ERCOT All", 700.7, 0.040, 0.006, 0.044),
    ("FRCC", "FRCC All", 763.7, 0.035, 0.004, 0.044),
    ("HIMS", "HICC Miscellaneous", 1087.7, 0.072, 0.012, 0.048),
    ("HIOA", "HICC Oahu", 1478.0, 0.136, 0.022, 0.048),
    ("MROE", "MRO East", 1297.3, 0.096, 0.014, 0.044),
    ("MROW", "MRO West", 862.6, 0.089, 0.013, 0.044),
    ("NEWE", "NPCC New England", 540.3, 0.059, 0.008, 0.044),
    ("NWPP", "WECC Northwest", 579.2, 0.047, 0.007, 0.044),
    ("NYCW", "NPCC NYC/Westchester", 883.2, 0.022, 0.002, 0.044),
    ("NYLI", "NPCC Long Island", 1170.6, 0.139, 0.018, 0.044),
    ("NYUP", "NPCC Upstate NY", 272.0, 0.011, 0.001, 0.044),
    ("PRMS", "Puerto Rico Miscellaneous", 1594.5, 0.085, 0.014, 0.140),
    ("RFCE", "RFC East", 618.6, 0.038, 0.005, 0.044),
    ("RFCM", "RFC Michigan", 982.5, 0.086, 0.012, 0.044),
    ("RFCW", "RFC West", 889.6, 0.068, 0.010, 0.044),
    ("RMPA", "WECC Rockies", 959.0, 0.079, 0.011, 0.044),
    ("SPNO", "SPP North", 779.0, 0.076, 0.011, 0.044),
    ("SPSO", "SPP South", 850.6, 0.051, 0.007, 0.044),
    ("SRMV", "SERC Mississippi Valley", 730.8, 0.029, 0.004, 0.044),
    ("SRMW", "SERC Midwest", 1150.4, 0.123, 0.018, 0.044),
    ("SRSO", "SERC South", 796.5, 0.056, 0.008, 0.044),
    ("SRTV", "SERC Tennessee Valley", 940.4, 0.080, 0.011, 0.044),
    ("SRVC", "SERC Virginia/Carolina", 637.4, 0.047, 0.007, 0.044),
    ("US", "US Average", 747.9, 0.054, 0.007, 0.044),
]

# Table 2 Mobile Combustion CO2 (kg CO2 per published unit).
MOBILE_CO2 = [
    ("Aviation Gasoline", 8.31, "gal (US)"),
    ("Biodiesel (100%)", 9.45, "gal (US)"),
    ("Compressed Natural Gas (CNG)", 0.05444, "scf"),
    ("Diesel Fuel", 10.21, "gal (US)"),
    ("Ethanol (100%)", 5.75, "gal (US)"),
    ("Kerosene-Type Jet Fuel", 9.75, "gal (US)"),
    ("Liquefied Natural Gas (LNG)", 4.50, "gal (US)"),
    ("Liquefied Petroleum Gases (LPG)", 5.68, "gal (US)"),
    ("Motor Gasoline", 8.78, "gal (US)"),
    ("Residual Fuel Oil", 11.27, "gal (US)"),
]

# Table 8 freight — kg CO2, g CH4, g N2O per published unit.
FREIGHT = [
    ("Medium- and Heavy-Duty Truck (vehicle-mile)", 1.298, 0.0115, 0.0376, "vehicle-mile"),
    ("Passenger Car (vehicle-mile)", 0.297, 0.0059, 0.0053, "vehicle-mile"),
    ("Light-Duty Truck (vehicle-mile)", 0.394, 0.0109, 0.0088, "vehicle-mile"),
    ("Medium- and Heavy-Duty Truck (short ton-mile)", 0.186, 0.0016, 0.0054, "short ton-mile"),
    ("Rail (short ton-mile)", 0.021, 0.0016, 0.0005, "short ton-mile"),
    ("Waterborne Craft (short ton-mile)", 0.077, 0.0310, 0.0020, "short ton-mile"),
    ("Aircraft (short ton-mile)", 1.086, 0.0, 0.0334, "short ton-mile"),
]

# Table 10 travel.
TRAVEL = [
    ("Passenger Car", 0.297, 0.0059, 0.0053, "vehicle-mile"),
    ("Light-Duty Truck", 0.394, 0.0109, 0.0088, "vehicle-mile"),
    ("Motorcycle", 0.368, 0.0888, 0.0188, "vehicle-mile"),
    ("Intercity Rail - Northeast Corridor", 0.023, 0.0010, 0.0001, "passenger-mile"),
    ("Intercity Rail - Other Routes", 0.149, 0.0120, 0.0040, "passenger-mile"),
    ("Intercity Rail - National Average", 0.096, 0.0080, 0.0020, "passenger-mile"),
    ("Commuter Rail", 0.133, 0.0105, 0.0026, "passenger-mile"),
    ("Transit Rail (subway, tram)", 0.093, 0.0075, 0.0010, "passenger-mile"),
    ("Bus", 0.066, 0.0046, 0.0019, "passenger-mile"),
    ("Air Travel - Short Haul (< 300 miles)", 0.215, 0.0126, 0.0081, "passenger-mile"),
    ("Air Travel - Medium Haul (300–2300 miles)", 0.120, 0.0006, 0.0056, "passenger-mile"),
    ("Air Travel - Long Haul (>= 2300 miles)", 0.143, 0.0006, 0.0078, "passenger-mile"),
]

# Table 9 WARM — MTCO2e per short ton. None = not published for that route.
WASTE_ROUTES = [
    "Recycled",
    "Landfilled",
    "Combusted",
    "Composted",
    "Anaerobically digested (dry)",
    "Anaerobically digested (wet)",
]
WASTE = [
    ("Aluminum Cans", 0.06, 0.02, 0.01, None, None, None),
    ("Aluminum Ingot", 0.04, 0.02, 0.01, None, None, None),
    ("Steel Cans", 0.32, 0.02, 0.01, None, None, None),
    ("Copper Wire", 0.18, 0.02, 0.01, None, None, None),
    ("Glass", 0.05, 0.02, 0.01, None, None, None),
    ("HDPE", 0.21, 0.02, 2.80, None, None, None),
    ("LDPE", None, 0.02, 2.80, None, None, None),
    ("PET", 0.23, 0.02, 2.05, None, None, None),
    ("LLDPE", None, 0.02, 2.80, None, None, None),
    ("PP", 0.20, 0.02, 2.80, None, None, None),
    ("PS", None, 0.02, 3.02, None, None, None),
    ("PVC", None, 0.02, 1.26, None, None, None),
    ("PLA", None, 0.02, 0.01, 0.13, None, None),
    ("Corrugated Containers", 0.11, 1.00, 0.05, None, None, None),
    ("Magazines/Third-class mail", 0.02, 0.46, 0.05, None, None, None),
    ("Newspaper", 0.02, 0.39, 0.05, None, None, None),
    ("Office Paper", 0.02, 1.41, 0.05, None, None, None),
    ("Phonebooks", 0.04, 0.39, 0.05, None, None, None),
    ("Textbooks", 0.04, 1.41, 0.05, None, None, None),
    ("Dimensional Lumber", None, 0.17, 0.05, None, None, None),
    ("Medium-density Fiberboard", None, 0.07, 0.05, None, None, None),
    ("Food Waste (non-meat)", None, 0.67, 0.05, 0.11, 0.14, 0.11),
    ("Food Waste (meat only)", None, 0.69, 0.05, 0.11, 0.14, 0.11),
    ("Beef", None, 0.64, 0.05, 0.11, 0.14, 0.11),
    ("Poultry", None, 0.73, 0.05, 0.11, 0.14, 0.11),
    ("Grains", None, 2.06, 0.05, 0.11, 0.14, 0.11),
    ("Bread", None, 1.49, 0.05, 0.11, 0.14, 0.11),
    ("Fruits and Vegetables", None, 0.28, 0.05, 0.11, 0.14, 0.11),
    ("Dairy Products", None, 0.72, 0.05, 0.11, 0.14, 0.11),
    ("Yard Trimmings", None, 0.36, 0.05, 0.14, 0.11, None),
    ("Grass", None, 0.28, 0.05, 0.14, 0.09, None),
    ("Leaves", None, 0.28, 0.05, 0.14, 0.12, None),
    ("Branches", None, 0.58, 0.05, 0.14, 0.15, None),
    ("Mixed Paper (general)", 0.07, 0.89, 0.05, None, None, None),
    ("Mixed Paper (primarily residential)", 0.07, 0.86, 0.05, None, None, None),
    ("Mixed Paper (primarily from offices)", 0.03, 0.84, 0.05, None, None, None),
    ("Mixed Metals", 0.23, 0.02, 0.01, None, None, None),
    ("Mixed Plastics", 0.22, 0.02, 2.34, None, None, None),
    ("Mixed Recyclables", 0.09, 0.75, 0.11, None, None, None),
    ("Food Waste", None, 0.68, 0.05, 0.11, None, None),
    ("Mixed Organics", None, 0.54, 0.05, 0.13, None, None),
    ("Mixed MSW", None, 0.58, 0.43, None, None, None),
    ("Carpet", None, 0.02, 1.68, None, None, None),
    ("Desktop CPUs", 0.01, 0.02, 0.40, None, None, None),
    ("Portable Electronic Devices", 0.02, 0.02, 0.89, None, None, None),
    ("Flat-panel Displays", 0.02, 0.02, 0.74, None, None, None),
    ("CRT Displays", None, 0.02, 0.64, None, None, None),
    ("Electronic Peripherals", 0.05, 0.02, 2.23, None, None, None),
    ("Hard-copy Devices", 0.01, 0.02, 1.92, None, None, None),
    ("Mixed Electronics", 0.02, 0.02, 0.96, None, None, None),
    ("Clay Bricks", None, 0.02, None, None, None, None),
    ("Concrete", 0.01, 0.02, None, None, None, None),
    ("Fly Ash", 0.01, 0.02, None, None, None, None),
    ("Tires", 0.10, 0.02, 2.21, None, None, None),
    ("Asphalt Concrete", 0.004, 0.02, None, None, None, None),
    ("Asphalt Shingles", 0.03, 0.02, 0.70, None, None, None),
    ("Drywall", None, 0.02, None, None, None, None),
    ("Fiberglass Insulation", 0.05, 0.02, None, None, None, None),
    ("Structural Steel", 0.04, 0.02, None, None, None, None),
    ("Vinyl Flooring", None, 0.02, 0.29, None, None, None),
    ("Wood Flooring", None, 0.18, 0.08, None, None, None),
]

# Table 11 AR6 GWPs (100-year).
GWP_PURE = [
    ("Carbon dioxide", "CO2", 1),
    ("Methane (non-fossil)", "CH4", 27),
    ("Fossil methane", "CH4", 29.8),
    ("Nitrous oxide", "N2O", 273),
    ("HFC-23", "CHF3", 14600),
    ("HFC-32", "CH2F2", 771),
    ("HFC-41", "CH3F", 135),
    ("HFC-125", "CHF2CF3", 3740),
    ("HFC-134", "CHF2CHF2", 1260),
    ("HFC-134a", "CH2FCF3", 1530),
    ("HFC-143", "CH2FCHF2", 364),
    ("HFC-143a", "CH3CF3", 5810),
    ("HFC-152", "CH2FCH2F", 21.5),
    ("HFC-152a", "CH3CHF2", 164),
    ("HFC-161", "CH3CH2F", 4.84),
    ("HFC-227ca", "CF3CF2CHF2", 2980),
    ("HFC-227ea", "CF3CHFCF3", 3600),
    ("HFC-236cb", "CH2FCF2CF3", 1350),
    ("HFC-236ea", "CHF2CHFCF3", 1500),
    ("HFC-236fa", "CF3CH2CF3", 8690),
    ("HFC-245ca", "CH2FCF2CHF2", 787),
    ("HFC-245cb", "CF3CF2CH3", 4550),
    ("HFC-245ea", "CHF2CHFCHF2", 255),
    ("HFC-245eb", "CH2FCHFCF3", 325),
    ("HFC-245fa", "CHF2CH2CF3", 962),
    ("HFC-263fb", "CH3CH2CF3", 74.8),
    ("HFC-272ca", "CH3CF2CH3", 599),
    ("HFC-329p", "CHF2CF2CF2CF3", 2890),
    ("HFC-365mfc", "CH3CF2CH2CF3", 914),
    ("HFC-43-10mee", "CF3CHFCHFCF2CF3", 1600),
    ("Sulfur hexafluoride", "SF6", 24300),
    ("Nitrogen trifluoride", "NF3", 17400),
    ("PFC-14", "CF4", 7380),
    ("PFC-116", "C2F6", 12400),
    ("PFC-218", "C3F8", 9290),
    ("PFC-318", "c-C4F8", 10200),
    ("PFC-31-10", "C4F10", 10000),
    ("PFC-41-12", "C5F12", 9220),
    ("PFC-51-14", "C6F14", 8620),
    ("PFC-61-16", "C7F16", 8410),
    ("PFC-71-18", "C8F18", 8260),
    ("PFC-91-18", "C10F18", 7480),
    ("PFC-1114", "C2F4", 0.004),
    ("PFC-1216", "C3F6", 0.09),
]

GWP_BLENDS = [
    ("R-401A", 21),
    ("R-401B", 18),
    ("R-401C", 25),
    ("R-402A", 2244),
    ("R-402B", 1421),
    ("R-403B", 3623),
    ("R-404A", 4728),
    ("R-406A", 0),
    ("R-407A", 2262),
    ("R-407B", 3001),
    ("R-407C", 1908),
    ("R-407D", 1748),
    ("R-408A", 2934),
    ("R-409A", 0),
    ("R-410A", 2256),
    ("R-410B", 2404),
    ("R-411A", 18),
    ("R-411B", 5),
    ("R-414A", 0),
    ("R-414B", 0),
    ("R-417A", 2508),
    ("R-422A", 3359),
    ("R-422B", 2700),
    ("R-422C", 3296),
    ("R-422D", 2917),
    ("R-424A", 2608),
    ("R-426A", 1614),
    ("R-427A", 2397),
    ("R-428A", 4061),
    ("R-434A", 3654),
    ("R-438A", 2425),
    ("R-441A", 0),
    ("R-442A", 2042),
    ("R-448A", 1494),
    ("R-449A", 1504),
    ("R-450A", 643),
    ("R-507A", 4775),
    ("R-508A", 13258),
    ("R-508B", 13412),
    ("R-513A", 673),
]


def slug(text: str) -> str:
    text = text.lower().replace("%", "pct").replace("<", "lt").replace(">", "gt")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")[:90]


def fmt(n: float) -> str:
    if abs(n) >= 1 and float(n) == int(n):
        return str(int(n))
    s = f"{n:.12f}".rstrip("0").rstrip(".")
    return s


def co2e(kg_co2: float, g_ch4: float, g_n2o: float) -> float:
    return kg_co2 + (g_ch4 / 1000.0) * GWP_CH4 + (g_n2o / 1000.0) * GWP_N2O


def parse_stationary() -> list[dict]:
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["Emission Factors Hub"]
    unit = "short ton"
    group = "Solid"
    rows = []
    seen = set()
    skip_names = {
        "Fuel Type",
        "Coal and Coke",
        "Other Fuels - Solid",
        "Biomass Fuels - Solid",
        "Natural Gas",
        "Other Fuels - Gaseous",
        "Biomass Fuels - Gaseous",
        "Petroleum Products",
        "Biomass Fuels - Liquid",
        "Source:",
        "Notes:",
    }
    for row in ws.iter_rows(min_row=14, max_row=89, values_only=True):
        name = row[2]
        d, h, i, j = row[3], row[7], row[8], row[9]
        if isinstance(d, str) and "mmBtu per" in d:
            if "short ton" in d:
                unit, group = "short ton", "Solid"
            elif "scf" in d:
                unit, group = "scf", "Gaseous"
            elif "gallon" in d:
                unit, group = "gal (US)", "Liquid"
            continue
        if not isinstance(name, str) or name in skip_names or name.startswith("Table"):
            # Natural Gas header vs data: data has numbers in h
            if name == "Natural Gas" and isinstance(h, (int, float)):
                pass
            else:
                continue
        if not isinstance(h, (int, float)):
            continue
        g_ch4 = float(i or 0)
        g_n2o = float(j or 0)
        kg = co2e(float(h), g_ch4, g_n2o)
        label = name.strip()
        if name == "Natural Gas":
            label = "Natural Gas"
            group = "Gaseous"
            unit = "scf"
        key = f"epa_stat_{slug(label)}_{slug(unit)}"
        if key in seen:
            continue
        seen.add(key)
        rows.append({"key": key, "name": label, "kg": kg, "unit": unit, "group": group})
    wb.close()
    return rows


def ts_row(fields: str) -> str:
    return "  { " + fields + " },"


def main() -> None:
    stationary = parse_stationary()
    lines = [
        "/** EPA Emission Factors Hub 2026 published types. Generated — do not edit by hand. */",
        "",
        "export const EPA_GWP_CH4_AR6 = 27",
        "export const EPA_GWP_N2O_AR6 = 273",
        "",
        "export type EpaEgridRow = { acronym: string; name: string; lbCo2: number; lbCh4: number; lbN2o: number; ggl: number }",
        "export type EpaFuelRow = { key: string; name: string; kg: number; unit: string; group: string; table: '1' | '2' }",
        "export type EpaGwpRow = { key: string; name: string; gwp: number; kind: 'pure' | 'blend' }",
        "export type EpaActivityRow = { key: string; name: string; kg: number; unit: string }",
        "export type EpaWasteRow = { key: string; material: string; route: string; kgPerShortTon: number }",
        "",
        "export const EPA_EGRID_SUBREGIONS: EpaEgridRow[] = [",
    ]
    for acronym, name, co2, ch4, n2o, ggl in EGRID:
        lines.append(
            ts_row(
                f"acronym: {acronym!r}, name: {name!r}, lbCo2: {fmt(co2)}, lbCh4: {fmt(ch4)}, lbN2o: {fmt(n2o)}, ggl: {fmt(ggl)}"
            )
        )
    lines.append("]")
    lines.append("")
    lines.append("export const EPA_STATIONARY_FUELS: EpaFuelRow[] = [")
    for row in stationary:
        lines.append(
            ts_row(
                f"key: {row['key']!r}, name: {row['name']!r}, kg: {fmt(row['kg'])}, unit: {row['unit']!r}, group: {row['group']!r}, table: '1'"
            )
        )
    lines.append("]")
    lines.append("")
    lines.append("export const EPA_MOBILE_CO2_FUELS: EpaFuelRow[] = [")
    aliases = {
        "Diesel Fuel": "diesel_us_gallon",
        "Motor Gasoline": "petrol_us_gallon",
    }
    for name, kg, unit in MOBILE_CO2:
        key = aliases.get(name, f"epa_mobile_{slug(name)}_{slug(unit)}")
        lines.append(
            ts_row(
                f"key: {key!r}, name: {name!r}, kg: {fmt(kg)}, unit: {unit!r}, group: 'Mobile CO2 (Table 2)', table: '2'"
            )
        )
    lines.append("]")
    lines.append("")
    lines.append("export const EPA_GWP_GASES: EpaGwpRow[] = [")
    aliases = {
        "HFC-134a": "r134a_epa_kg",
        "Fossil methane": "mine_ch4_epa_fossil_kg",
    }
    for name, _formula, gwp in GWP_PURE:
        key = aliases.get(name, f"epa_gwp_{slug(name)}_kg")
        lines.append(ts_row(f"key: {key!r}, name: {name!r}, gwp: {fmt(gwp)}, kind: 'pure'"))
    blend_alias = {"R-404A": "r404a_epa_kg", "R-410A": "r410a_epa_kg"}
    for name, gwp in GWP_BLENDS:
        key = blend_alias.get(name, f"epa_gwp_{slug(name)}_kg")
        lines.append(ts_row(f"key: {key!r}, name: {name!r}, gwp: {fmt(gwp)}, kind: 'blend'"))
    lines.append("]")
    lines.append("")
    lines.append("export const EPA_FREIGHT_ROWS: EpaActivityRow[] = [")
    for name, kg_co2, g_ch4, g_n2o, unit in FREIGHT:
        kg = co2e(kg_co2, g_ch4, g_n2o)
        lines.append(
            ts_row(f"key: {'epa_freight_' + slug(name)!r}, name: {name!r}, kg: {fmt(kg)}, unit: {unit!r}")
        )
    lines.append("]")
    lines.append("")
    lines.append("export const EPA_TRAVEL_ROWS: EpaActivityRow[] = [")
    for name, kg_co2, g_ch4, g_n2o, unit in TRAVEL:
        kg = co2e(kg_co2, g_ch4, g_n2o)
        lines.append(
            ts_row(f"key: {'epa_travel_' + slug(name)!r}, name: {name!r}, kg: {fmt(kg)}, unit: {unit!r}")
        )
    lines.append("]")
    lines.append("")
    lines.append("export const EPA_WASTE_ROWS: EpaWasteRow[] = [")
    for material, *vals in WASTE:
        for route, val in zip(WASTE_ROUTES, vals):
            if val is None:
                continue
            kg = float(val) * 1000  # MTCO2e / short ton → kg CO2e / short ton
            key = f"epa_waste_{slug(material)}_{slug(route)}"
            lines.append(
                ts_row(
                    f"key: {key!r}, material: {material!r}, route: {route!r}, kgPerShortTon: {fmt(kg)}"
                )
            )
    lines.append("]")
    lines.append("")
    OUT.write_text("\n".join(lines) + "\n")
    print(
        f"wrote EPA hub data: eGRID {len(EGRID)}, stationary {len(stationary)}, "
        f"mobile {len(MOBILE_CO2)}, GWP {len(GWP_PURE)+len(GWP_BLENDS)}, "
        f"freight {len(FREIGHT)}, travel {len(TRAVEL)}, waste materials {len(WASTE)}"
    )


if __name__ == "__main__":
    main()
