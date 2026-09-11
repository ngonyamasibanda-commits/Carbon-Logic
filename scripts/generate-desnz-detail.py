#!/usr/bin/env python3
"""Extract DESNZ 2026 conversion rows into src/lib/desnz-detail-data.ts.

Reads the official full-set workbook from the uploads folder. The xlsx is not
committed; re-run this script when DESNZ publishes a new year.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path(
    "/home/ubuntu/.cursor/projects/workspace/uploads/ghg-conversion-factors-2026-full-set_86e9.xlsx"
)
OUT = Path("/workspace/src/lib/desnz-detail-data.ts")

# Existing catalogue keys we must keep stable (tests + saved entries).
ALIASES: dict[tuple, str] = {}


def slug(text: str) -> str:
    text = text.lower().replace("%", "pct").replace(">", "gt").replace("×", "x")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")[:80]


def fmt_num(n: float) -> str:
    if n == int(n) and abs(n) >= 1:
        return str(int(n))
    s = f"{n:.10f}".rstrip("0").rstrip(".")
    return s


def cell_num(ws, r, c):
    v = ws.cell(r, c).value
    return float(v) if isinstance(v, (int, float)) else None


def fill_down(ws, start, end, cols=(1, 2)):
    last = {c: None for c in cols}
    rows = []
    for r in range(start, end + 1):
        for c in cols:
            v = ws.cell(r, c).value
            if v not in (None, ""):
                last[c] = str(v).strip()
        rows.append((r, dict(last)))
    return rows


rows_out: list[dict] = []
lookup: dict[str, str] = {}
seen_keys: set[str] = set()


def add_row(
    key: str,
    name: str,
    kg: float | None,
    unit: str,
    scope: str,
    category: str,
    path: str,
    *,
    wtt_key: str | None = None,
    ev_key: str | None = None,
    ev_td_key: str | None = None,
    method: str | None = None,
    alias: str | None = None,
):
    if kg is None:
        return
    if unit in ("miles",):
        lookup[path] = alias or key
        return
    final_key = alias or key
    lookup[path] = final_key
    if final_key in seen_keys:
        # Same published row reused (km tables shared delivery/freight).
        return
    seen_keys.add(final_key)
    rec = {
        "key": final_key,
        "name": name,
        "kg": kg,
        "unit": unit,
        "scope": scope,
        "category": category,
    }
    if wtt_key:
        rec["wttKey"] = wtt_key
    if ev_key:
        rec["evKey"] = ev_key
    if ev_td_key:
        rec["evTdKey"] = ev_td_key
    if method:
        rec["method"] = method
    rows_out.append(rec)


print("Loading workbook…")
wb = load_workbook(XLSX, data_only=True)

VAN_FUELS = [
    (4, "Diesel"),
    (8, "Petrol"),
    (12, "CNG"),
    (16, "LPG"),
    (20, "Unknown"),
    (24, "Plug-in Hybrid Electric Vehicle"),
    (28, "Battery Electric Vehicle"),
]
CAR_FUELS = [
    (4, "Diesel"),
    (8, "Petrol"),
    (12, "Hybrid"),
    (16, "CNG"),
    (20, "LPG"),
    (24, "Unknown"),
    (28, "Plug-in Hybrid Electric Vehicle"),
    (32, "Battery Electric Vehicle"),
]
LADEN = [
    (4, "Average laden"),
    (8, "0% Laden"),
    (12, "50% Laden"),
    (16, "100% Laden"),
]
WTT_VAN_FUELS = [
    (4, "Diesel"),
    (5, "Petrol"),
    (6, "CNG"),
    (7, "LPG"),
    (8, "Unknown"),
    (9, "Plug-in Hybrid Electric Vehicle"),
    (10, "Battery Electric Vehicle"),
]
WTT_LADEN = [
    (4, "Average laden"),
    (5, "0% Laden"),
    (6, "50% Laden"),
    (7, "100% Laden"),
]
EV_FUELS = [
    (4, "Plug-in Hybrid Electric Vehicle"),
    (8, "Battery Electric Vehicle"),
]


def wtt_map(ws, start, end, activity_prefix: str, fuels, unit_ok):
    """path -> kg for WTT sheet (one column per fuel or laden)."""
    out = {}
    last_a = last_b = None
    for r in range(start, end + 1):
        a, b, u = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if a:
            last_a = str(a).strip()
        if b:
            last_b = str(b).strip()
        if u not in unit_ok or not last_b:
            continue
        for col, fuel in fuels:
            kg = cell_num(ws, r, col)
            if kg is None:
                continue
            out[f"{last_b}|{fuel}|{u}"] = kg
    return out


# ── WTT delivery / freight (for wttKey pairing) ────────────────────────────
wtt_ws = wb["WTT- delivery vehs & freight"]
wtt_van = {}
last_b = None
for r in range(20, 35):
    b, u = wtt_ws.cell(r, 2).value, wtt_ws.cell(r, 3).value
    if b:
        last_b = str(b).strip()
    if u not in ("km", "tonne.km") or not last_b:
        continue
    for col, fuel in WTT_VAN_FUELS:
        kg = cell_num(wtt_ws, r, col)
        if kg is None:
            continue
        wtt_van[f"{last_b}|{fuel}|{u}"] = kg

wtt_hgv = {}
last_a = last_b = None
for r in range(36, 90):
    a, b, u = wtt_ws.cell(r, 1).value, wtt_ws.cell(r, 2).value, wtt_ws.cell(r, 3).value
    if a:
        last_a = str(a).strip()
    if b:
        last_b = str(b).strip()
    if u not in ("km", "tonne.km") or not last_b:
        continue
    refrigerated = "refrigerated" in (last_a or "").lower() and "non-refrigerated" not in (
        last_a or ""
    )
    for col, laden in WTT_LADEN:
        kg = cell_num(wtt_ws, r, col)
        if kg is None:
            continue
        wtt_hgv[f"{last_b}|{refrigerated}|{laden}|{u}"] = kg


def wtt_van_key(cls, fuel, unit):
    path = f"wtt_van|{cls}|{fuel}|{unit}"
    if path in lookup:
        return lookup[path]
    kg = wtt_van.get(f"{cls}|{fuel}|{unit}")
    if kg is None:
        return None
    key = f"wtt_van_{slug(cls)}_{slug(fuel)}_{'tkm' if unit == 'tonne.km' else unit}"
    add_row(
        key,
        f"WTT van — {cls}, {fuel}",
        kg,
        "tkm" if unit == "tonne.km" else unit,
        "Scope 3",
        "WTT freight",
        path,
    )
    return key


def wtt_hgv_key(cls, refrigerated, laden, unit):
    path = f"wtt_hgv|{cls}|{int(refrigerated)}|{laden}|{unit}"
    if path in lookup:
        return lookup[path]
    kg = wtt_hgv.get(f"{cls}|{refrigerated}|{laden}|{unit}")
    if kg is None:
        return None
    body = "refrigerated" if refrigerated else "non_ref"
    key = f"wtt_hgv_{body}_{slug(cls)}_{slug(laden)}_{'tkm' if unit == 'tonne.km' else unit}"
    add_row(
        key,
        f"WTT HGV — {cls}, {'refrigerated' if refrigerated else 'non-refrigerated'}, {laden}",
        kg,
        "tkm" if unit == "tonne.km" else unit,
        "Scope 3",
        "WTT freight",
        path,
    )
    return key


# ── UK electricity for EVs / T&D for EVs ───────────────────────────────────
def extract_ev_table(sheet_name: str, key_prefix: str, category: str, scope: str):
    ws = wb[sheet_name]
    last_a = last_b = None
    for r in range(24, min(ws.max_row, 120) + 1):
        a, b, u = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if a:
            last_a = str(a).strip()
        if b:
            last_b = str(b).strip()
        if u not in ("km", "miles", "tonne.km") or not last_b:
            continue
        for col, fuel in EV_FUELS:
            kg = cell_num(ws, r, col)
            if kg is None:
                continue
            unit = "tkm" if u == "tonne.km" else u
            key = f"{key_prefix}_{slug(last_a)}_{slug(last_b)}_{slug(fuel)}_{unit}"
            path = f"{key_prefix}|{last_a}|{last_b}|{fuel}|{u}"
            add_row(
                key,
                f"{category} — {last_b}, {fuel}",
                kg,
                unit,
                scope,
                category,
                path,
            )


extract_ev_table("UK electricity for EVs", "ev_elec", "UK electricity for EVs", "Scope 2")
extract_ev_table("UK electricity T&D for EVs", "ev_td", "UK electricity T&D for EVs", "Scope 3")


def ev_keys(activity: str, typ: str, fuel: str, unit: str):
    ev = lookup.get(f"ev_elec|{activity}|{typ}|{fuel}|{unit}")
    td = lookup.get(f"ev_td|{activity}|{typ}|{fuel}|{unit}")
    return ev, td


# ── Delivery vehicles: vans (km) + HGV (km × laden) ────────────────────────
dv = wb["Delivery vehicles"]
last_a = last_b = None
for r in range(24, 36):
    a, b, u = dv.cell(r, 1).value, dv.cell(r, 2).value, dv.cell(r, 3).value
    if a:
        last_a = str(a).strip()
    if b:
        last_b = str(b).strip()
    if u != "km" or last_a != "Vans" or not last_b:
        continue
    for col, fuel in VAN_FUELS:
        kg = cell_num(dv, r, col)
        if kg is None:
            continue
        key = f"van_{slug(last_b)}_{slug(fuel)}_km"
        path = f"van|km|{last_b}|{fuel}"
        wtt = wtt_van_key(last_b, fuel, "km")
        evk = td = None
        if "Electric" in fuel:
            evk, td = ev_keys("Vans", last_b, fuel, "km")
        add_row(
            key,
            f"Van {last_b}, {fuel} (vehicle-km)",
            kg,
            "km",
            "Scope 1",
            "Fleet vehicles",
            path,
            wtt_key=wtt,
            ev_key=evk,
            ev_td_key=td,
        )

last_a = last_b = None
for r in range(36, 80):
    a, b, u = dv.cell(r, 1).value, dv.cell(r, 2).value, dv.cell(r, 3).value
    if a:
        last_a = str(a).strip()
    if b:
        last_b = str(b).strip()
    if u != "km" or not last_a or not last_b or "HGV" not in last_a:
        continue
    refrigerated = "refrigerated" in last_a.lower() and "non-refrigerated" not in last_a.lower()
    for col, laden in LADEN:
        kg = cell_num(dv, r, col)
        if kg is None:
            continue
        body = "ref" if refrigerated else "nr"
        key = f"hgv_{body}_{slug(last_b)}_{slug(laden)}_km"
        path = f"hgv|km|{last_b}|{int(refrigerated)}|{laden}"
        wtt = wtt_hgv_key(last_b, refrigerated, laden, "km")
        add_row(
            key,
            f"HGV {last_b}, {'refrigerated' if refrigerated else 'non-refrigerated'}, {laden} (vehicle-km)",
            kg,
            "km",
            "Scope 1",
            "Fleet vehicles",
            path,
            wtt_key=wtt,
        )


# ── Freighting goods: vans tkm/km (km already added), HGV tkm, air, rail, sea
fg = wb["Freighting goods"]
last_a = last_b = None
for r in range(24, 96):
    a, b, u = fg.cell(r, 1).value, fg.cell(r, 2).value, fg.cell(r, 3).value
    if a:
        last_a = str(a).strip()
    if b:
        last_b = str(b).strip()
    if not last_b or u not in ("km", "tonne.km"):
        continue
    if last_a == "Vans":
        for col, fuel in VAN_FUELS:
            kg = cell_num(fg, r, col)
            if kg is None:
                continue
            unit = "tkm" if u == "tonne.km" else "km"
            key = f"van_{slug(last_b)}_{slug(fuel)}_{unit}"
            path = f"van|{unit}|{last_b}|{fuel}"
            alias = None
            if last_b.startswith("Average") and fuel == "Diesel" and unit == "tkm":
                alias = "freight_van_tkm"
            wtt = wtt_van_key(last_b, fuel, u)
            evk = td = None
            if "Electric" in fuel:
                evk, td = ev_keys("Vans", last_b, fuel, u)
            add_row(
                key,
                f"Van freight {last_b}, {fuel}",
                kg,
                unit,
                "Scope 3",
                "Road freight",
                path,
                wtt_key=wtt,
                ev_key=evk,
                ev_td_key=td,
                alias=alias,
            )
    elif last_a and "HGV" in last_a:
        refrigerated = "refrigerated" in last_a.lower() and "non-refrigerated" not in last_a.lower()
        for col, laden in LADEN:
            kg = cell_num(fg, r, col)
            if kg is None:
                continue
            unit = "tkm" if u == "tonne.km" else "km"
            body = "ref" if refrigerated else "nr"
            key = f"hgv_{body}_{slug(last_b)}_{slug(laden)}_{unit}"
            path = f"hgv|{unit}|{last_b}|{int(refrigerated)}|{laden}"
            alias = None
            if unit == "tkm" and laden == "Average laden" and not refrigerated:
                if last_b == "Average non-refrigerated HGVs":
                    alias = "freight_road_tkm"
                elif last_b == "Average non-refrigerated rigids":
                    alias = "freight_road_rigid_tkm"
                elif last_b == "Average non-refrigerated artics":
                    alias = "freight_road_artic_tkm"
            wtt = wtt_hgv_key(last_b, refrigerated, laden, u)
            add_row(
                key,
                f"HGV freight {last_b}, {'refrigerated' if refrigerated else 'non-refrigerated'}, {laden}",
                kg,
                unit,
                "Scope 3",
                "Road freight",
                path,
                wtt_key=wtt,
                alias=alias,
            )

# Air freight: col4 with RF, col8 without RF
last_b = None
for r in range(96, 105):
    b, u = fg.cell(r, 2).value, fg.cell(r, 3).value
    if b:
        last_b = str(b).strip()
    if u != "tonne.km" or not last_b:
        continue
    with_rf = cell_num(fg, r, 4)
    no_rf = cell_num(fg, r, 8)
    haul = last_b.split(",")[0].strip()
    for rf_label, kg, suffix in (
        ("With RF (DESNZ default)", with_rf, "rf"),
        ("Without RF", no_rf, "no_rf"),
    ):
        key = f"freight_air_{slug(haul)}_{suffix}_tkm"
        path = f"air|{haul}|{rf_label}"
        alias = None
        if haul == "Short-haul" and suffix == "rf":
            alias = "freight_air_tkm"
        if haul == "Short-haul" and suffix == "no_rf":
            alias = "freight_air_tkm_no_rf"
        add_row(
            key,
            f"Air freight {last_b}, {rf_label}",
            kg,
            "tkm",
            "Scope 3",
            "Air freight",
            path,
            alias=alias,
        )

# Rail
for r in range(105, 110):
    b, u = fg.cell(r, 2).value, fg.cell(r, 3).value
    if str(u) == "tonne.km" and b:
        add_row(
            "freight_rail_tkm",
            "Rail freight",
            cell_num(fg, r, 4),
            "tkm",
            "Scope 3",
            "Rail freight",
            "rail|Freight train",
            alias="freight_rail_tkm",
        )

# Sea: Activity / Type / Size / Unit / kg
last_a = last_b = last_size = None
for r in range(108, 200):
    a, b, c, d = (
        fg.cell(r, 1).value,
        fg.cell(r, 2).value,
        fg.cell(r, 3).value,
        fg.cell(r, 4).value,
    )
    if a:
        last_a = str(a).strip()
    if b:
        last_b = str(b).strip()
    # Size is col3, unit col4 for sea tables
    if last_a not in ("Sea tanker", "Cargo ship"):
        continue
    size = str(c).strip() if c not in (None, "Size", "Unit") else None
    unit = d
    kg = cell_num(fg, r, 5)
    if size in (None, "Activity", "Type") or unit not in ("tonne.km",):
        # some rows might put unit in col3
        continue
    if kg is None:
        continue
    key = f"sea_{slug(last_b)}_{slug(size)}_tkm"
    path = f"sea|{last_b}|{size}"
    alias = None
    if last_b == "Container ship" and size.lower() == "average":
        alias = "freight_sea_tkm"
    if last_b == "Bulk carrier" and size.lower() == "average":
        alias = "freight_sea_bulk_tkm"
    if last_b.lower().startswith("roro") and size.lower() == "average":
        alias = "freight_sea_roro_tkm"
    add_row(
        key,
        f"{last_b} ({size})",
        kg,
        "tkm",
        "Scope 3",
        "Sea freight",
        path,
        alias=alias,
    )


# ── Passenger vehicles (company cars, km) ──────────────────────────────────
pv = wb["Passenger vehicles"]
last_a = last_b = None
for r in range(24, 70):
    a, b, u = pv.cell(r, 1).value, pv.cell(r, 2).value, pv.cell(r, 3).value
    if a:
        last_a = str(a).strip()
    if b:
        last_b = str(b).strip()
    if u != "km" or not last_b:
        continue
    table = "car_segment" if last_a == "Cars (by market segment)" else (
        "car_size" if last_a == "Cars (by size)" else ("motorbike" if last_a == "Motorbike" else None)
    )
    if not table:
        continue
    fuels = CAR_FUELS if table.startswith("car") else [(4, "Petrol")]
    for col, fuel in fuels:
        kg = cell_num(pv, r, col)
        if kg is None:
            continue
        key = f"{table}_{slug(last_b)}_{slug(fuel)}_km"
        path = f"{table}|km|{last_b}|{fuel}"
        alias = None
        if table == "car_size" and last_b == "Average car" and fuel == "Diesel":
            alias = "car_diesel_km"
        if table == "car_size" and last_b == "Average car" and fuel == "Petrol":
            alias = "car_petrol_km"
        if table == "car_size" and last_b == "Average car" and fuel == "Unknown":
            alias = "commute_car_km"
        if table == "motorbike" and last_b == "Average":
            alias = "commute_motorbike_km"
        evk = td = None
        if "Electric" in fuel:
            activity = "Cars (by size)" if table == "car_size" else last_a
            evk, td = ev_keys(activity, last_b, fuel, "km")
        add_row(
            key,
            f"{last_b}, {fuel}",
            kg,
            "km",
            "Scope 1",
            "Passenger vehicles",
            path,
            ev_key=evk,
            ev_td_key=td,
            alias=alias,
        )
        # Scope 3 duplicate path for business travel / commuting
        lookup[f"s3|{path}"] = lookup[path]


# ── Business travel land (taxis, bus, rail) ────────────────────────────────
bl = wb["Business travel- land"]
last_a = last_b = None
for r in range(68, 95):
    a, b, u = bl.cell(r, 1).value, bl.cell(r, 2).value, bl.cell(r, 3).value
    if a:
        last_a = str(a).strip()
    if b:
        last_b = str(b).strip()
    if not last_b or u not in ("passenger.km", "km"):
        continue
    if last_a not in ("Taxis", "Bus", "Rail"):
        continue
    kg = cell_num(bl, r, 4)
    unit = "pkm" if u == "passenger.km" else "km"
    key = f"{slug(last_a)}_{slug(last_b)}_{unit}"
    path = f"land|{last_a}|{last_b}|{unit}"
    alias = None
    if last_b == "Regular taxi" and unit == "pkm":
        alias = "taxi_pkm"
    if last_b == "Regular taxi" and unit == "km":
        alias = "taxi_km"
    if last_b == "Average local bus" and unit == "pkm":
        alias = "commute_bus_pkm"
    if last_b == "National rail" and unit == "pkm":
        alias = "commute_rail_pkm"
    add_row(
        key,
        f"{last_b}",
        kg,
        unit,
        "Scope 3",
        "Business travel",
        path,
        alias=alias,
    )

# Ferry
sea_t = wb["Business travel- sea"]
last_b = None
for r in range(16, 24):
    b, u = sea_t.cell(r, 2).value, sea_t.cell(r, 3).value
    if b:
        last_b = str(b).strip()
    if u != "passenger.km" or not last_b:
        continue
    kg = cell_num(sea_t, r, 4)
    key = f"ferry_{slug(last_b)}_pkm"
    add_row(
        key,
        f"Ferry — {last_b}",
        kg,
        "pkm",
        "Scope 3",
        "Business travel",
        f"ferry|{last_b}",
    )

# Flights: col5 with RF, col9 without; fill haul/class
air = wb["Business travel- air"]
last_haul = last_cls = None
for r in range(22, 38):
    haul, cls, u = air.cell(r, 2).value, air.cell(r, 3).value, air.cell(r, 4).value
    if haul:
        last_haul = str(haul).strip()
    if cls:
        last_cls = str(cls).strip()
    if u != "passenger.km" or not last_haul or not last_cls:
        continue
    with_rf = cell_num(air, r, 5)
    no_rf = cell_num(air, r, 9)
    haul_short = last_haul.split(",")[0].strip()
    for rf_label, kg, suffix in (
        ("With RF (DESNZ default)", with_rf, "rf"),
        ("Without RF", no_rf, "no_rf"),
    ):
        key = f"flight_{slug(haul_short)}_{slug(last_cls)}_{suffix}_pkm"
        path = f"flight|{haul_short}|{last_cls}|{rf_label}"
        alias = None
        if suffix == "rf":
            if haul_short == "Domestic" and last_cls == "Average passenger":
                alias = "flight_domestic_pkm"
            if haul_short == "Short-haul" and last_cls == "Economy class":
                alias = "flight_shorthaul_pkm"
            if haul_short == "Long-haul" and last_cls == "Economy class":
                alias = "flight_longhaul_economy_pkm"
            if haul_short == "Long-haul" and last_cls == "Business class":
                alias = "flight_longhaul_business_pkm"
        add_row(
            key,
            f"{last_haul}, {last_cls}, {rf_label}",
            kg,
            "pkm",
            "Scope 3",
            "Business travel",
            path,
            alias=alias,
        )


# ── Waste ──────────────────────────────────────────────────────────────────
waste = wb["Waste disposal"]
ROUTE_COLS = [
    (4, "Re-use"),
    (5, "Open-loop"),
    (6, "Closed-loop"),
    (7, "Combustion"),
    (8, "Composting"),
    (9, "Landfill"),
    (10, "Anaerobic digestion"),
]
last_a = last_b = None
for r in range(24, 120):
    a, b, u = waste.cell(r, 1).value, waste.cell(r, 2).value, waste.cell(r, 3).value
    if a:
        last_a = str(a).strip()
    if b:
        last_b = str(b).strip()
    if str(u) not in ("tonnes", "tonne") or not last_b or last_b == "Waste type":
        continue
    for col, route in ROUTE_COLS:
        kg = cell_num(waste, r, col)
        if kg is None:
            continue
        key = f"waste_{slug(last_b)}_{slug(route)}_t"
        path = f"waste|{last_b}|{route}"
        alias = None
        if last_b in ("Aggregates", "Asphalt", "Bricks", "Concrete") and route == "Landfill":
            alias = "waste_landfill_kg"
        if last_b in ("Aggregates", "Average construction", "Asphalt", "Concrete") and route in (
            "Closed-loop",
            "Open-loop",
        ):
            alias = "waste_recycling_kg"
        if last_b == "Average construction" and route == "Combustion":
            alias = "waste_combustion_kg"
        add_row(
            key,
            f"{last_b} — {route}",
            kg,
            "t",
            "Scope 3",
            "Waste",
            path,
            alias=alias,
        )


# ── Refrigerants ───────────────────────────────────────────────────────────
ref = wb["Refrigerant & other"]
last_a = None
for r in range(18, 200):
    a, b, u = ref.cell(r, 1).value, ref.cell(r, 2).value, ref.cell(r, 3).value
    if a:
        last_a = str(a).strip()
    if str(u) != "kg" or not b:
        continue
    kg = cell_num(ref, r, 4)
    if kg is None:
        continue
    name = str(b).strip()
    key = f"ref_{slug(name)}_kg"
    path = f"ref|{name}"
    alias = None
    if name == "Carbon dioxide":
        alias = "co2_kg"
    if name == "HFC-134a":
        alias = "r134a_kg"
    if name == "R404A":
        alias = "r404a_kg"
    if name == "R410A":
        alias = "r410a_kg"
    add_row(
        key,
        name,
        kg,
        "kg",
        "Scope 1",
        "Refrigerants",
        path,
        method="gwp",
        alias=alias,
    )


# ── Fuels + WTT-fuels ──────────────────────────────────────────────────────
FUEL_UNIT = {
    "litres": "L",
    "tonnes": "t",
    "cubic metres": "m³",
    "kWh (Net CV)": "kWh",
    "kWh (Gross CV)": "kWh",
    "kWh": "kWh",
    "kg": "kg",
    "GJ": "GJ",
    "million litres": "ML",
}

wtt_fuel_kg: dict[str, float] = {}
wtt_f = wb["WTT- fuels"]
last_b = None
for r in range(21, 160):
    b, u = wtt_f.cell(r, 2).value, wtt_f.cell(r, 3).value
    if b:
        last_b = str(b).strip()
    kg = cell_num(wtt_f, r, 4)
    if last_b and u and kg is not None:
        wtt_fuel_kg[f"{last_b}|{u}"] = kg

wtt_bio_kg: dict[str, float] = {}
wtt_b = wb["WTT- bioenergy"]
last_b = None
for r in range(18, 90):
    b, u = wtt_b.cell(r, 2).value, wtt_b.cell(r, 3).value
    if b:
        last_b = str(b).strip()
    kg = cell_num(wtt_b, r, 4)
    if last_b and u and kg is not None:
        wtt_bio_kg[f"{last_b}|{u}"] = kg


def add_wtt_energy(prefix: str, name: str, unit: str, kg: float | None, category: str):
    if kg is None:
        return None
    key = f"wtt_{prefix}_{slug(name)}_{slug(unit)}"
    path = f"wtt_{prefix}|{name}|{unit}"
    alias = None
    if prefix == "fuel" and name == "Diesel (average biofuel blend)" and unit == "litres":
        alias = "wtt_diesel_litre"
    if prefix == "fuel" and name == "Petrol (average biofuel blend)" and unit == "litres":
        alias = "wtt_petrol_litre"
    if prefix == "fuel" and name == "LPG" and unit == "litres":
        alias = "wtt_lpg_litre"
    if prefix == "fuel" and name == "Gas oil" and unit == "litres":
        alias = "wtt_gas_oil_litre"
    if prefix == "fuel" and name == "Natural gas" and unit == "kWh (Gross CV)":
        alias = "wtt_natural_gas_kwh"
    add_row(
        key,
        f"WTT — {name} ({unit})",
        kg,
        FUEL_UNIT.get(str(unit), str(unit)),
        "Scope 3",
        category,
        path,
        alias=alias,
    )
    return alias or key


fuels = wb["Fuels"]
last_a = last_b = None
for r in range(22, 160):
    a, b, u = fuels.cell(r, 1).value, fuels.cell(r, 2).value, fuels.cell(r, 3).value
    if a:
        last_a = str(a).strip()
    if b:
        last_b = str(b).strip()
    kg = cell_num(fuels, r, 4)
    if not last_b or not u or kg is None:
        continue
    if last_a not in ("Gaseous fuels", "Liquid fuels", "Solid fuels"):
        continue
    unit = str(u).strip()
    key = f"fuel_{slug(last_b)}_{slug(unit)}"
    path = f"fuel|{last_b}|{unit}"
    alias = None
    if last_b == "Diesel (average biofuel blend)" and unit == "litres":
        alias = "diesel_litre"
    if last_b == "Petrol (average biofuel blend)" and unit == "litres":
        alias = "petrol_litre"
    if last_b == "LPG" and unit == "litres":
        alias = "lpg_litre"
    if last_b == "Gas oil" and unit == "litres":
        alias = "gas_oil_litre"
    if last_b == "Natural gas" and unit == "cubic metres":
        alias = "natural_gas_m3"
    if last_b == "Natural gas" and unit == "kWh (Gross CV)":
        alias = "natural_gas_kwh"
    if last_b == "Marine gas oil" and unit == "litres":
        alias = "marine_gas_oil_litre"
    if last_b == "Aviation turbine fuel" and unit == "litres":
        alias = "aviation_turbine_litre"
    wtt = add_wtt_energy("fuel", last_b, unit, wtt_fuel_kg.get(f"{last_b}|{unit}"), "WTT fuels")
    add_row(
        key,
        f"{last_b} ({unit})",
        kg,
        FUEL_UNIT.get(unit, unit),
        "Scope 1",
        "Site fuel",
        path,
        wtt_key=wtt,
        alias=alias,
    )

bio = wb["Bioenergy"]
last_a = last_b = None
for r in range(19, 85):
    a, b, u = bio.cell(r, 1).value, bio.cell(r, 2).value, bio.cell(r, 3).value
    if a:
        last_a = str(a).strip()
    if b:
        last_b = str(b).strip()
    kg = cell_num(bio, r, 4)
    if not last_b or not u or kg is None:
        continue
    if last_a not in ("Biofuel", "Biomass", "Biogas"):
        continue
    unit = str(u).strip()
    key = f"bio_{slug(last_b)}_{slug(unit)}"
    path = f"bio|{last_b}|{unit}"
    alias = "hvo_litre" if last_b == "Biodiesel HVO" and unit == "litres" else None
    wtt = add_wtt_energy("bio", last_b, unit, wtt_bio_kg.get(f"{last_b}|{unit}"), "WTT bioenergy")
    add_row(
        key,
        f"{last_b} ({unit})",
        kg,
        FUEL_UNIT.get(unit, unit),
        "Scope 1",
        "Bioenergy",
        path,
        wtt_key=wtt,
        alias=alias,
    )

# ── Material use (primary / reused / recycled) ─────────────────────────────
MAT_COLS = [
    (4, "Primary material production"),
    (5, "Re-used"),
    (6, "Open-loop source"),
    (7, "Closed-loop source"),
]
mat = wb["Material use"]
last_b = None
for r in range(22, 95):
    b, u = mat.cell(r, 2).value, mat.cell(r, 3).value
    if b:
        last_b = str(b).strip()
    if str(u) != "tonnes" or not last_b:
        continue
    for col, origin in MAT_COLS:
        kg = cell_num(mat, r, col)
        if kg is None:
            continue
        key = f"material_{slug(last_b)}_{slug(origin)}_t"
        path = f"material|{last_b}|{origin}"
        alias = None
        if origin == "Primary material production":
            alias = {
                "Aggregates": "material_aggregates_t",
                "Asphalt": "material_asphalt_t",
                "Bricks": "material_bricks_t",
                "Concrete": "material_concrete_t",
                "Insulation": "material_insulation_t",
                "Metals": "material_metals_t",
                "Plasterboard": "material_plasterboard_t",
                "Wood": "material_timber_t",
                "Glass": "material_glass_t",
                "Plastics: PVC (incl. forming)": "material_pvc_t",
            }.get(last_b)
        add_row(
            key,
            f"{last_b} — {origin}",
            kg,
            "t",
            "Scope 3",
            "Bulk materials",
            path,
            alias=alias,
        )

# ── Homeworking ────────────────────────────────────────────────────────────
hw = wb["Homeworking"]
for r in range(22, 25):
    name = hw.cell(r, 1).value
    unit = hw.cell(r, 2).value
    kg = cell_num(hw, r, 3)
    if not name or kg is None:
        continue
    key = f"homeworking_{slug(str(name))}_hour"
    add_row(
        key,
        str(name),
        kg,
        "hour",
        "Scope 3",
        "Homeworking",
        f"homeworking|{name}",
    )

# Also alias crew keys
if "land|Bus|Average local bus|pkm" in lookup:
    lookup["crew|Shuttle bus"] = lookup["land|Bus|Average local bus|pkm"]
if "land|Rail|National rail|pkm" in lookup:
    lookup["crew|Rail"] = lookup["land|Rail|National rail|pkm"]
# crew van historically used average car unknown
if "car_size|km|Average car|Unknown" in lookup:
    lookup["crew|Crew van"] = lookup["car_size|km|Average car|Unknown"]
    # existing key
    lookup["crew|Crew van"] = "crew_van_km"


print(f"rows={len(rows_out)} lookup={len(lookup)} unique_keys={len(seen_keys)}")

# Sanity: HGV rigid 3.5-7.5 avg km
p = "hgv|km|Rigid (>3.5 - 7.5 tonnes)|0|Average laden"
print("HGV 3.5-7.5 km key", lookup.get(p))
row = next((x for x in rows_out if x["key"] == lookup.get(p)), None)
print("  kg", row and row["kg"])
print("van class I diesel km", lookup.get("van|km|Class I (up to 1.305 tonnes)|Diesel"))
print("sea container avg", lookup.get("sea|Container ship|Average"))
print("flight short economy rf", lookup.get("flight|Short-haul|Economy class|With RF (DESNZ default)"))
print("waste soils landfill", lookup.get("waste|Soils|Landfill"))
print("ref R32", lookup.get("ref|HFC-32"))

# Collect option lists from lookup paths
van_classes = sorted({p.split("|")[2] for p in lookup if p.startswith("van|km|")})
hgv_classes = sorted({p.split("|")[2] for p in lookup if p.startswith("hgv|tkm|")})
sea_types = sorted({p.split("|")[1] for p in lookup if p.startswith("sea|")})
waste_types = sorted({p.split("|")[1] for p in lookup if p.startswith("waste|")})
ref_names = [p.split("|", 1)[1] for p in lookup if p.startswith("ref|")]
car_sizes = sorted({p.split("|")[2] for p in lookup if p.startswith("car_size|km|")})
car_segments = sorted({p.split("|")[2] for p in lookup if p.startswith("car_segment|km|")})
print("van_classes", van_classes)
print("hgv_classes", hgv_classes)
print("sea_types", sea_types)
print("waste n", len(waste_types), "ref n", len(ref_names))
print("car_sizes", car_sizes)

# Write TS
SKIP = {
    "freight_van_tkm",
    "freight_road_tkm",
    "freight_road_rigid_tkm",
    "freight_road_artic_tkm",
    "freight_air_tkm",
    "freight_air_tkm_no_rf",
    "freight_rail_tkm",
    "freight_sea_tkm",
    "freight_sea_bulk_tkm",
    "freight_sea_roro_tkm",
    "waste_landfill_kg",
    "waste_recycling_kg",
    "waste_combustion_kg",
    "car_diesel_km",
    "car_petrol_km",
    "commute_car_km",
    "commute_motorbike_km",
    "commute_bus_pkm",
    "commute_rail_pkm",
    "taxi_pkm",
    "taxi_km",
    "flight_domestic_pkm",
    "flight_shorthaul_pkm",
    "flight_longhaul_economy_pkm",
    "flight_longhaul_business_pkm",
    "r134a_kg",
    "r404a_kg",
    "r410a_kg",
    "co2_kg",
    "crew_van_km",
    "crew_bus_pkm",
    "crew_rail_pkm",
    "diesel_litre",
    "petrol_litre",
    "lpg_litre",
    "gas_oil_litre",
    "natural_gas_m3",
    "natural_gas_kwh",
    "marine_gas_oil_litre",
    "aviation_turbine_litre",
    "hvo_litre",
    "wtt_diesel_litre",
    "wtt_petrol_litre",
    "wtt_lpg_litre",
    "wtt_gas_oil_litre",
    "wtt_natural_gas_kwh",
    "material_aggregates_t",
    "material_asphalt_t",
    "material_bricks_t",
    "material_concrete_t",
    "material_insulation_t",
    "material_metals_t",
    "material_plasterboard_t",
    "material_timber_t",
    "material_glass_t",
    "material_pvc_t",
}

new_rows = [r for r in rows_out if r["key"] not in SKIP]

def ts_str(s: str) -> str:
    return json.dumps(s, ensure_ascii=False)

lines = []
lines.append("/**")
lines.append(" * DESNZ / UK GHG Conversion Factors 2026 — detailed rows generated")
lines.append(" * from the official full-set workbook (not committed).")
lines.append(" * Re-run: python3 scripts/generate-desnz-detail.py")
lines.append(" */")
lines.append("")
lines.append("export type DesnzDetailRow = {")
lines.append("  key: string")
lines.append("  name: string")
lines.append("  kg: number")
lines.append("  unit: string")
lines.append("  scope: 'Scope 1' | 'Scope 2' | 'Scope 3'")
lines.append("  category: string")
lines.append("  wttKey?: string")
lines.append("  evKey?: string")
lines.append("  evTdKey?: string")
lines.append("  method?: 'gwp'")
lines.append("}")
lines.append("")
lines.append("export const DESNZ_DETAIL_ROWS: DesnzDetailRow[] = [")
for rec in new_rows:
    parts = [
        f"key: {ts_str(rec['key'])}",
        f"name: {ts_str(rec['name'])}",
        f"kg: {fmt_num(rec['kg'])}",
        f"unit: {ts_str(rec['unit'])}",
        f"scope: {ts_str(rec['scope'])}",
        f"category: {ts_str(rec['category'])}",
    ]
    if rec.get("wttKey"):
        parts.append(f"wttKey: {ts_str(rec['wttKey'])}")
    if rec.get("evKey"):
        parts.append(f"evKey: {ts_str(rec['evKey'])}")
    if rec.get("evTdKey"):
        parts.append(f"evTdKey: {ts_str(rec['evTdKey'])}")
    if rec.get("method"):
        parts.append(f"method: {ts_str(rec['method'])}")
    lines.append("  { " + ", ".join(parts) + " },")
lines.append("]")
lines.append("")
lines.append("export const DESNZ_LOOKUP: Record<string, string> = {")
for path, key in sorted(lookup.items()):
    lines.append(f"  {ts_str(path)}: {ts_str(key)},")
lines.append("}")
lines.append("")

OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
print("wrote", OUT, "bytes", OUT.stat().st_size, "new_rows", len(new_rows))
